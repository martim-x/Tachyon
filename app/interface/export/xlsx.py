"""Адаптер экспорта отчёта в xlsx через openpyxl (docs/03 §S7.1 T2).

Адаптер только форматирует данные домена: агрегация уже посчитана в
reporting (NFR-3), здесь нет pandas и запросов. Деньги пишутся
Decimal-значениями в числовые ячейки с форматом MONEY_CELL_FORMAT -
никаких float для денег и никаких текстовых сумм. Пустой период даёт
файл только с русскими шапками и нулями сводки (S7.1).
"""

from collections.abc import Sequence
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.domains.reporting.schemas import (
    ExportOperationDTO,
    ExportReportDTO,
    SeriesPointDTO,
)
from app.interface.export.base import ExportFile

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MONEY_CELL_FORMAT = "#,##0.00"
PERCENT_CELL_FORMAT = "0.00"
DATE_CELL_FORMAT = "yyyy-mm-dd"

KIND_LABELS = {
    "income": "Доход",
    "expense": "Расход",
}

HEADER_FONT = Font(bold=True)

OPERATIONS_HEADERS = ("Дата", "Категория", "Тип", "Сумма")
OPERATIONS_WIDTHS = (14, 30, 10, 16)

CATEGORIES_HEADERS = (
    "Категория",
    "Сумма",
    "Доля, %",
    "Лимит",
    "Остаток",
    "Использовано, %",
    "Превышен",
)
CATEGORIES_WIDTHS = (30, 16, 12, 14, 14, 18, 12)

SUMMARY_WIDTHS = (24, 20)

COLUMN_LETTERS = ("A", "B", "C", "D", "E", "F", "G")


def _set_widths(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for column_letter, width in zip(COLUMN_LETTERS, widths, strict=False):
        sheet.column_dimensions[column_letter].width = width


def _write_headers(
    sheet: Worksheet,
    headers: tuple[str, ...],
    widths: tuple[int, ...],
) -> None:
    """Шапка листа русскими названиями в жирном начертании + ширины колонок."""

    for index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = HEADER_FONT

    _set_widths(sheet, widths)


def _set_text(
    sheet: Worksheet,
    row: int,
    column: int,
    text: str,
) -> None:
    """Пользовательский текст пишется строкой, а не формулой (CWE-1236).

    openpyxl объявляет значением формулу любую строку, начинающуюся с «=»:
    категория имени «=HYPERLINK(...)» ушла бы в файл живой формулой и
    выполнилась при открытии выгрузки. data_type = «s» принудительно
    сохраняет исходный текст.
    """

    cell = sheet.cell(row=row, column=column, value=text)
    if cell.data_type != "s":
        cell.data_type = "s"


def _set_money(
    sheet: Worksheet,
    row: int,
    column: int,
    amount: Decimal | None,
) -> None:
    """Числовая денежная ячейка с двумя знаками, а не текст."""

    if amount is None:
        return

    cell = sheet.cell(row=row, column=column, value=amount)
    cell.number_format = MONEY_CELL_FORMAT


def _set_percent(
    sheet: Worksheet,
    row: int,
    column: int,
    value: Decimal | None,
) -> None:
    if value is None:
        return

    cell = sheet.cell(row=row, column=column, value=value)
    cell.number_format = PERCENT_CELL_FORMAT


def _fill_operations_sheet(
    sheet: Worksheet,
    operations: Sequence[ExportOperationDTO],
) -> None:
    _write_headers(sheet, OPERATIONS_HEADERS, OPERATIONS_WIDTHS)

    for offset, operation in enumerate(operations, start=1):
        row = offset + 1

        date_cell = sheet.cell(row=row, column=1, value=operation.occurred_on)
        date_cell.number_format = DATE_CELL_FORMAT

        _set_text(sheet, row, 2, operation.category_name)
        sheet.cell(row=row, column=3, value=KIND_LABELS[operation.kind])
        _set_money(sheet, row, 4, operation.amount)


def _fill_categories_sheet(
    sheet: Worksheet,
    categories: Sequence[SeriesPointDTO],
) -> None:
    _write_headers(sheet, CATEGORIES_HEADERS, CATEGORIES_WIDTHS)

    for offset, category in enumerate(categories, start=1):
        row = offset + 1

        _set_text(sheet, row, 1, category.category_name)
        _set_money(sheet, row, 2, category.amount)
        _set_percent(sheet, row, 3, category.share_percent)
        _set_money(sheet, row, 4, category.limit)
        _set_money(sheet, row, 5, category.limit_remaining)
        _set_percent(sheet, row, 6, category.limit_percent)
        sheet.cell(row=row, column=7, value=category.limit_exceeded)


def _fill_summary_sheet(sheet: Worksheet, report: ExportReportDTO) -> None:
    _set_widths(sheet, SUMMARY_WIDTHS)

    summary: tuple[tuple[str, date | Decimal], ...] = (
        ("Период с", report.date_from),
        ("Период по", report.date_to),
        ("Доходы", report.income),
        ("Расходы", report.expense),
        ("Баланс", report.balance),
    )

    for row_offset, (label, value) in enumerate(summary, start=1):
        label_cell = sheet.cell(row=row_offset, column=1, value=label)
        label_cell.font = HEADER_FONT

        cell = sheet.cell(row=row_offset, column=2, value=value)
        if isinstance(value, Decimal):
            cell.number_format = MONEY_CELL_FORMAT
        else:
            cell.number_format = DATE_CELL_FORMAT


def _filename(report: ExportReportDTO) -> str:
    """Имя файла включает период, чтобы отчёты не перетирали друг друга."""

    return f"report-{report.date_from:%Y-%m-%d}_{report.date_to:%Y-%m-%d}.xlsx"


class XlsxReportExporter:
    """xlsx-реализация порта ReportExporter (docs/02 §export)."""

    def export(self, report: ExportReportDTO) -> ExportFile:
        workbook = Workbook()

        operations_sheet = workbook.worksheets[0]
        operations_sheet.title = "Операции"
        _fill_operations_sheet(operations_sheet, report.operations)

        _fill_categories_sheet(workbook.create_sheet("По категориям"), report.categories)
        _fill_summary_sheet(workbook.create_sheet("Сводка"), report)

        buffer = BytesIO()
        workbook.save(buffer)

        return ExportFile(
            filename=_filename(report),
            media_type=XLSX_MEDIA_TYPE,
            content=buffer.getvalue(),
        )
