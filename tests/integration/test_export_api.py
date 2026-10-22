"""Integration-тесты GET /api/v1/reports/monthly/export (docs/03 §S7.1).

Имена тестов содержат 'export' для выборки `pytest -k export` и
'isolation' для `pytest -k isolation` (docs/06 §Структура тестов).
Чужие операции пользователя B не попадают в файл пользователя A.
"""

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from tests.conftest import CategoryFactory, LoggedInClient, OperationFactory

pytestmark = pytest.mark.integration

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PERIOD_PARAMS = {"from": "2026-10-01", "to": "2026-10-31"}


def _workbook(response_content: bytes) -> Workbook:
    return load_workbook(BytesIO(response_content))


def _reloaded(number: object) -> Decimal:
    """Точное сравнение числа, перечитанного из xlsx (int/float)."""

    return Decimal(str(number))


def _operations_rows(response_content: bytes) -> list[tuple[object, str, str, object]]:
    """Строки листа «Операции» без шапки: дата, категория, тип, сумма."""

    sheet = _workbook(response_content)["Операции"]
    return [
        (row[0].value, str(row[1].value), str(row[2].value), row[3].value)
        for row in sheet.iter_rows(min_row=2, max_col=4)
    ]


def _category_names(response_content: bytes) -> list[str]:
    sheet = _workbook(response_content)["По категориям"]
    return [str(row[0].value) for row in sheet.iter_rows(min_row=2, max_col=1)]


async def _export_excel(client: AsyncClient, **extra: str):
    params = dict(PERIOD_PARAMS, format="excel", **extra)
    return await client.get("/api/v1/reports/monthly/export", params=params)


async def test_monthly_export_excel_returns_file_with_headers(
    api_client_a: LoggedInClient,
) -> None:
    """S7.1: MIME-тип, Content-Disposition с периодом, три листа, русские шапки."""

    response = await _export_excel(api_client_a.client)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    assert response.headers["content-disposition"] == (
        'attachment; filename="report-2026-10-01_2026-10-31.xlsx"'
    )

    workbook = _workbook(response.content)
    assert workbook.sheetnames == ["Операции", "По категориям", "Сводка"]
    assert [cell.value for cell in workbook["Операции"][1]] == [
        "Дата",
        "Категория",
        "Тип",
        "Сумма",
    ]
    assert workbook["Операции"]["A1"].font.bold is True


async def test_monthly_export_excel_contains_period_operations_and_summary(
    api_client_a: LoggedInClient,
    category_factory: CategoryFactory,
    operation_factory: OperationFactory,
) -> None:
    """S7.1: лист «Операции» содержит постаточно свои операции, «Сводка» - итоги."""

    food = await category_factory(api_client_a.user, name="Продукты", kind="expense")
    salary = await category_factory(api_client_a.user, name="Зарплата", kind="income")

    await operation_factory(
        api_client_a.user,
        amount=Decimal("2000.00"),
        category=salary,
        kind="income",
        occurred_at=datetime(2026, 10, 1, 12, tzinfo=UTC),
    )
    await operation_factory(
        api_client_a.user,
        amount=Decimal("100.00"),
        category=food,
        kind="expense",
        occurred_at=datetime(2026, 10, 5, 12, tzinfo=UTC),
    )
    # Операция до периода не попадает в выгрузку периода октябрь.
    await operation_factory(
        api_client_a.user,
        amount=Decimal("77.00"),
        category=food,
        kind="expense",
        occurred_at=datetime(2026, 9, 5, 12, tzinfo=UTC),
    )

    response = await _export_excel(api_client_a.client)
    assert response.status_code == 200

    rows = _operations_rows(response.content)
    assert len(rows) == 2
    assert rows[0][1] == "Зарплата"
    assert rows[0][2] == "Доход"
    assert _reloaded(rows[0][3]) == Decimal("2000.00")
    assert rows[1][1] == "Продукты"
    assert rows[1][2] == "Расход"
    assert _reloaded(rows[1][3]) == Decimal("100.00")

    workbook = _workbook(response.content)
    summary = workbook["Сводка"]
    assert summary["A1"].value == "Период с"
    assert summary["A3"].value == "Доходы"
    assert summary["B3"].value == 2000
    assert summary["B4"].value == 100
    assert summary["B5"].value == 1900
    assert summary["B3"].number_format == "#,##0.00"

    # Денежная ячейка листа «Операции» отдаётся числом, а не текстом.
    money_cell = workbook["Операции"].cell(row=2, column=4)
    assert not isinstance(money_cell.value, str)
    assert money_cell.number_format == "#,##0.00"


async def test_monthly_export_excel_empty_period_writes_headers_only(
    api_client_a: LoggedInClient,
) -> None:
    """S7.1: у пользователя без операций - файл с шапками и нулевой сводкой."""

    response = await _export_excel(api_client_a.client)

    assert response.status_code == 200

    workbook = _workbook(response.content)
    operations = workbook["Операции"]
    assert operations.max_row == 1
    assert operations["A1"].value == "Дата"
    assert _workbook(response.content)["По категориям"].max_row == 1

    summary = workbook["Сводка"]
    assert summary["B3"].value == 0
    assert summary["B4"].value == 0
    assert summary["B5"].value == 0


async def test_monthly_export_excel_endpoint_isolation_ignores_other_user_operations(
    api_client_a: LoggedInClient,
    api_client_b: LoggedInClient,
    category_factory: CategoryFactory,
    operation_factory: OperationFactory,
) -> None:
    """S7.1 T4: чужие операции в выгрузку пользователя A не попадают.

    Модель - test_forecast_endpoint_isolation_*; подстановка чужого
    user_id в query игнорируется, цифры в файле только свои.
    """

    food = await category_factory(api_client_a.user, name="Продукты", kind="expense")
    salary = await category_factory(api_client_a.user, name="Зарплата", kind="income")

    await operation_factory(
        api_client_a.user,
        amount=Decimal("100.00"),
        category=food,
        kind="expense",
        occurred_at=datetime(2026, 10, 5, 12, tzinfo=UTC),
    )
    await operation_factory(
        api_client_a.user,
        amount=Decimal("2000.00"),
        category=salary,
        kind="income",
        occurred_at=datetime(2026, 10, 1, 12, tzinfo=UTC),
    )

    foreign = await category_factory(
        api_client_b.user,
        name="Чужие расходы",
        kind="expense",
    )
    await operation_factory(
        api_client_b.user,
        amount=Decimal("9999.00"),
        category=foreign,
        kind="expense",
        occurred_at=datetime(2026, 10, 5, 12, tzinfo=UTC),
    )

    # Чужой user_id в query игнорируется: фильтр берётся из UserContext.
    response = await _export_excel(
        api_client_a.client,
        user_id=str(api_client_b.user.id),
    )
    assert response.status_code == 200

    rows = _operations_rows(response.content)
    assert len(rows) == 2
    assert "Чужие расходы" not in _category_names(response.content)
    assert all(_reloaded(row[3]) != Decimal("9999.00") for row in rows)
    assert summary_totals(response.content) == (
        Decimal("2000.00"),
        Decimal("100.00"),
    )


async def test_monthly_export_excel_rejects_reversed_period(
    api_client_a: LoggedInClient,
) -> None:
    response = await _export_excel(
        api_client_a.client,
        **{"from": "2026-10-31", "to": "2026-10-01"},
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_ERROR"


async def test_monthly_export_excel_rejects_unknown_format(
    api_client_a: LoggedInClient,
) -> None:
    response = await api_client_a.client.get(
        "/api/v1/reports/monthly/export",
        params=dict(PERIOD_PARAMS, format="gif"),
    )

    assert response.status_code == 422


async def test_monthly_export_excel_requires_authentication() -> None:
    """Без Bearer-токена выгрузка не отдаётся (401, не 500)."""

    from app.interface.api.app import create_app

    application = create_app()

    async with AsyncClient(
        transport=ASGITransport(app=application),
        base_url="http://test",
    ) as client:
        response = await client.get(
            "/api/v1/reports/monthly/export",
            params=PERIOD_PARAMS,
        )

    assert response.status_code == 401


def summary_totals(response_content: bytes) -> tuple[Decimal, Decimal]:
    """Доходы и расходы листа «Сводка» для точного сравнения."""

    sheet = _workbook(response_content)["Сводка"]
    return _reloaded(sheet["B3"].value), _reloaded(sheet["B4"].value)
